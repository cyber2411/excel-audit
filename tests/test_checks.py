"""Smoke tests for the four checks (stdlib unittest, no extra deps)."""

import unittest

import openpyxl

from excel_audit.checks import (
    audit_workbook,
    count_volatile_functions,
    find_broken_refs,
    find_inconsistent_columns,
)


def _workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "=B1*C1"
    ws["A2"] = "=B2*C2"
    ws["A3"] = "=B3+C3"  # inconsistent with the column pattern
    ws["B1"] = "='No Such Sheet'!A1"
    ws["C1"] = "=TODAY()"
    return wb


class TestChecks(unittest.TestCase):
    def test_broken_refs(self):
        issues = find_broken_refs(_workbook())
        self.assertTrue(any("No Such Sheet" in i["reason"] for i in issues))

    def test_inconsistent_column(self):
        issues = find_inconsistent_columns(_workbook())
        self.assertEqual(len(issues), 1)
        self.assertIn("column A", issues[0]["cell"])

    def test_volatile_count(self):
        counts = count_volatile_functions(_workbook())
        self.assertEqual(counts.get("TODAY"), 1)

    def test_audit_totals(self):
        result = audit_workbook(_workbook())
        self.assertEqual(result["issue_total"], 2)
        self.assertEqual(result["volatile_total"], 1)


if __name__ == "__main__":
    unittest.main()
